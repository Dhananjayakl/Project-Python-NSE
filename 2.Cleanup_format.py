# This script processes CSV files from a specified input directory, cleans the data, and saves it as Excel files in an output directory.
#             file_path = os.path.join(data_dir, filename)
import pandas as pd
import os
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

# Define directories
input_folder = r"C:\Users\dhana\Desktop\python_stock_test\Project-Python-NSE\nse_data"
output_folder = r"C:\Users\dhana\Desktop\python_stock_test\Project-Python-NSE\output"

# Create output folder if it doesn't exist
os.makedirs(output_folder, exist_ok=True)

# Process all CSV files
for file in os.listdir(input_folder):
    if file.endswith(".csv"):
        file_path = os.path.join(input_folder, file)
        symbol = os.path.splitext(file)[0]

        # Read the CSV
        df = pd.read_csv(file_path, skiprows=2)
        df.columns = ["Date", "Close", "High", "Low", "Open", "Volume", "Symbol"]

        # Round values to 2 decimal places
        for col in ["Close", "High", "Low", "Open"]:
            df[col] = pd.to_numeric(df[col], errors="coerce").round(2)

        # Output Excel path
        output_excel_path = os.path.join(output_folder, f"{symbol}.xlsx")

        # Write to Excel
        df.to_excel(output_excel_path, index=False)

        # Adjust column widths
        wb = load_workbook(output_excel_path)
        ws = wb.active
        for col in ws.columns:
            max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max_length + 2
        wb.save(output_excel_path)

print(" All files processed and saved in 'output' folder.")
